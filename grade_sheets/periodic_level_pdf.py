import os
import logging
import uuid
from django.conf import settings
from docx import Document
from docx2pdf import convert
import pythoncom
from grade_sheets.pdf_utils import replace_placeholders
from grade_sheets.helpers import get_grade_sheet_data
from students.helper import get_students_by_level
from PyPDF2 import PdfMerger

logger = logging.getLogger(__name__)

def generate_periodic_level_pdf(template_path, level_id, academic_year_id):
    """Generate a single PDF for all students in a level, with each student on a separate page."""
    try:
        if not os.path.exists(template_path):
            logger.error(f"Template not found: {template_path}")
            raise FileNotFoundError(f"Template not found: {template_path}")

        output_dir = os.path.join(settings.MEDIA_ROOT, 'output_gradesheets')
        temp_dir = os.path.join(settings.MEDIA_ROOT, 'temp')
        os.makedirs(output_dir, exist_ok=True)
        os.makedirs(temp_dir, exist_ok=True)
        logger.info(f"Ensured directories exist: {output_dir}, {temp_dir}")

        # Verify directory permissions
        try:
            test_file = os.path.join(temp_dir, f"test_{uuid.uuid4()}.txt")
            with open(test_file, 'w') as f:
                f.write('test')
            os.remove(test_file)
            logger.info(f"Successfully tested write permissions in {temp_dir}")
        except PermissionError as e:
            logger.error(f"Permission error in {temp_dir}: {str(e)}")
            return []

        # Get all students for the level
        students = get_students_by_level(level_id)
        if not students:
            logger.warning(f"No students found for level_id={level_id}")
            return []

        # Initialize PDF merger and temporary file list
        merger = PdfMerger()
        temp_files = []
        docx_paths = []
        pdf_paths = []

        # Generate all DOCX files
        for student in students:
            # Get student data
            data = get_grade_sheet_data(student_id=student.id, level_id=level_id, academic_year_id=academic_year_id)
            if not data or 'name' not in data:
                logger.warning(f"No data found for student_id={student.id}, level_id={level_id}, academic_year_id={academic_year_id}")
                continue

            # Generate individual student DOCX using replace_placeholders
            doc = Document(template_path)
            doc = replace_placeholders(doc, data)
            safe_name = data.get('name', 'student').replace(' ', '_').replace(':', '_').replace('/', '_').replace('\\', '_')
            temp_filename = f"temp_report_card_{safe_name}_{student.id}_{academic_year_id}_{uuid.uuid4().hex}.docx"
            temp_path = os.path.join(temp_dir, temp_filename)
            doc.save(temp_path)
            docx_paths.append(temp_path)
            temp_files.append(temp_path)
            logger.info(f"Saved temporary student DOCX: {temp_path}")

        # Batch convert all DOCX files to PDF
        pythoncom.CoInitialize()
        try:
            for docx_path in docx_paths:
                temp_pdf_filename = docx_path.replace('.docx', '.pdf')
                convert(docx_path, temp_pdf_filename)
                if not os.path.exists(temp_pdf_filename):
                    logger.error(f"PDF not created at {temp_pdf_filename}")
                    continue
                pdf_paths.append(temp_pdf_filename)
                temp_files.append(temp_pdf_filename)
                logger.info(f"Generated temporary PDF: {temp_pdf_filename}")
                # Append to merger
                merger.append(temp_pdf_filename)
        finally:
            pythoncom.CoUninitialize()

        # Save the combined PDF
        pdf_filename = f"level_{level_id}_{academic_year_id}.pdf"
        pdf_path = os.path.join(output_dir, pdf_filename)
        merger.write(pdf_path)
        merger.close()
        logger.info(f"PDF generated for level {level_id}: {pdf_path}")

        # Clean up temporary files
        try:
            for temp_file in temp_files:
                if os.path.exists(temp_file):
                    os.remove(temp_file)
                    logger.info(f"Cleaned up temporary file: {temp_file}")
        except Exception as e:
            logger.warning(f"Failed to clean up temporary files: {str(e)}")

        if os.path.exists(pdf_path):
            return [pdf_path]
        else:
            logger.error(f"PDF not created at {pdf_path}")
            return []

    except Exception as e:
        logger.error(f"Error generating level grade PDF: {str(e)}", exc_info=True)
        return []