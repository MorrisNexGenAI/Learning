from docx import Document
import os
import logging
from django.conf import settings
from grade_sheets.helpers import get_grade_sheet_data
from docx2pdf import convert

logger = logging.getLogger(__name__)

def generate_student_grade_pdf(template_path, student_id, level_id, academic_year_id):
    """Generate PDF for a single student with name and grades replaced."""
    try:
        if not os.path.exists(template_path):
            logger.error(f"Template not found: {template_path}")
            raise FileNotFoundError(f"Template not found: {template_path}")

        output_dir = os.path.join(settings.MEDIA_ROOT, 'output_gradesheets')
        temp_dir = os.path.join(settings.MEDIA_ROOT, 'temp')
        os.makedirs(output_dir, exist_ok=True)
        os.makedirs(temp_dir, exist_ok=True)

        # Get student data
        data = get_grade_sheet_data(student_id=student_id, level_id=level_id, academic_year_id=academic_year_id)
        if not data or 'name' not in data:
            logger.warning(f"No data found for student_id={student_id}, level_id={level_id}, academic_year_id={academic_year_id}")
            return []

        # Use replace_placeholders from periodic_pdf
        from grade_sheets.periodic_pdf import replace_placeholders
        doc = Document(template_path)
        doc = replace_placeholders(doc, data)
        safe_name = data.get('name', 'student').replace(' ', '_').replace(':', '_').replace('/', '_').replace('\\', '_')
        temp_filename = f"temp_report_card_{safe_name}_{academic_year_id}.docx"
        temp_path = os.path.join(temp_dir, temp_filename)
        doc.save(temp_path)
        logger.info(f"Saved temporary DOCX: {temp_path}")

        pdf_filename = temp_filename.replace('.docx', '.pdf')
        pdf_path = os.path.join(output_dir, pdf_filename)
        convert(temp_path, pdf_path)
        if os.path.exists(pdf_path):
            logger.info(f"PDF generated: {pdf_path}")
            return [pdf_path]
        else:
            logger.error(f"PDF not created at {pdf_path}")
            return []
    except Exception as e:
        logger.error(f"Error generating student grade PDF: {str(e)}")
        return []